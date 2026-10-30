"""吊架批次結果 → Excel 報表（mep_tray/hanger_report.py）與網頁 ④ 區「匯出 Excel」。

測試資料是依 HangerPlacer.cs 的 HangerRunReport 欄位手工組成的合成批次，不是實際 Revit 輸出。
"""
import copy
import io
import json
import threading

import pytest
from openpyxl import load_workbook

from mep_tray import acad
from mep_tray import export_dxf as X
from mep_tray import hanger_report as HR
from mep_tray import webui as W


def item(sid, tray, station, status="new", reason="spacing", fitting=None, **kw):
    it = {"SupportId": sid, "TrayId": tray, "TrayUniqueId": f"uid-{fitting or tray}", "FittingId": fitting,
          "TrayHeightMm": 100.0, "StationMm": station, "TrayPointMm": [1000.0 + station, 2000.0, 3000.0],
          "BearingPointMm": [1000.0 + station, 2000.0, 2950.0], "InsertionPointMm": [1000.0 + station, 2000.0, 2900.0],
          "RotationDeg": 90.0, "TrayDirectionDeg": 0.0, "CrossMemberLengthMm": 400.0, "RodLengthMm": None,
          "HitElementId": None, "HitLinkInstanceId": None, "HitCategory": None, "HitZMm": None,
          "Reason": reason, "Status": status, "Detail": "", "ElementId": None, "TypeName": None}
    it.update(kw)
    return it


def batch():
    items = [
        item("H001", "101", 300.0),
        item("H002", "101", 1800.0, status="exists", Detail="本工具已放置（元素 9001）"),
        item("H003", "101", 3300.0, status="pending", Detail="附近有非本工具（手動）吊架（元素 9002）"),
        item("H004", "101", 2550.0, reason="gap_fix"),
        item("H005", "102", 150.0, reason="near_fitting", status="created", ElementId="7001"),
        item("H006", "102", 1650.0, status="failed", Detail="InvalidOperationException: x"),
        item("H007", "201", 785.398, reason="elbow_mid", fitting="201"),
        item("H008", "202", 0.0, reason="tee_mid", fitting="202"),
    ]
    return {"BatchId": "20261021-125357-abcdef12", "Entry": "selection", "RodMode": "none",
            "Coordinates": "Revit 內部座標（mm）", "Outcome": "預覽", "PreflightErrors": [],
            "LayoutIssues": [{"Level": "pending", "Code": "fixed_remainder", "PieceId": "102",
                              "LocationMm": [1.0, 2.0, 3.0], "Detail": "剩餘 12.000 mm"}],
            "Items": items, "RevitWarnings": [], "Notes": ["n1"], "TestFailNth": None, "InstanceParams": []}


def workbook(data: bytes):
    return load_workbook(io.BytesIO(data))


def sheet_rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


# ───────────── 欄位對照表 ─────────────
def test_shipped_column_map_is_valid_and_has_no_invented_parameter_names():
    cols = HR.load_columns()
    assert all(c["revit_param"] is None for c in cols), "實際參數名稱尚未取得，不得預填"
    keys = [c["key"] for c in cols]
    for k in ("support_id", "tray_id", "type", "spacing_mm", "status", "tray_x_mm", "tray_y_mm", "tray_z_mm"):
        assert k in keys


@pytest.mark.parametrize("mutate, msg", [
    (lambda o: o.update(version=2), "version"),
    (lambda o: o["columns"][1].update(revit_param=" 名稱"), "revit_param"),
    (lambda o: o["columns"][1].update(revit_param=""), "revit_param"),
    (lambda o: (o["columns"][1].update(revit_param="P"), o["columns"][2].update(revit_param="P")), "重複"),
    (lambda o: o["columns"][1].update(source=None), "衍生"),
    (lambda o: o["columns"][1].update(source="Station Mm"), "source"),
    (lambda o: o["columns"][1].update(extra=1), "只能"),
    (lambda o: o["columns"][2].update(key="support_id"), "key"),
])
def test_column_map_validation_names_the_problem(mutate, msg):
    o = json.loads(HR.COLUMNS_PATH.read_text(encoding="utf-8"))
    mutate(o)
    with pytest.raises(HR.HangerReportError, match=msg):
        HR.validate_columns(o)


def test_filled_parameter_name_appears_in_the_mapping_sheet(tmp_path):
    o = json.loads(HR.COLUMNS_PATH.read_text(encoding="utf-8"))
    o["columns"][1]["revit_param"] = "吊架編號參數"
    p = tmp_path / "cols.json"
    p.write_text(json.dumps(o, ensure_ascii=False), encoding="utf-8")
    wb = workbook(HR.build_workbook(HR.parse_batch(batch()), HR.load_columns(p)))
    m = sheet_rows(wb["參數對照"])
    assert m[0] == ["報表欄位", "欄位代碼", "來源（批次 JSON）", "Revit 吊架參數名稱"]
    assert m[2] == ["吊架編號", "support_id", "SupportId", "吊架編號參數"]
    assert m[3][3] == "未設定"


# ───────────── 衍生欄位 ─────────────
def test_types_statuses_and_spacing():
    b = HR.parse_batch(batch())
    assert [HR.hanger_type(i) for i in b["Items"]] == ["直線段"] * 6 + ["彎頭中點", "三通中點"]
    # 101：300 → 1800 → 2550 → 3300（依 station 排序，與列順序無關）；102：150 → 1650；中點吊架留空
    assert HR.spacings(b["Items"]) == [None, 1500.0, 750.0, 750.0, None, 1500.0, None, None]
    assert HR.counts(b) == {"new": 4, "created": 1, "exists": 1, "pending": 1, "failed": 1}


@pytest.mark.parametrize("reason, fitting", [("elbow_mid", None), ("spacing", "301"), ("", None), ("other", None)])
def test_inconsistent_or_unknown_reason_is_not_guessed(reason, fitting):
    t = HR.hanger_type(HR.parse_batch({"BatchId": "b", "Items": [item("H1", "1", 0.0, reason=reason, fitting=fitting)]})["Items"][0])
    assert t.startswith("無法判定")


# ───────────── 輸入驗證 ─────────────
@pytest.mark.parametrize("mutate, field", [
    (lambda o: o.pop("BatchId"), "BatchId"),
    (lambda o: o.update(Items={}), "Items"),
    (lambda o: o["Items"][2].update(Status="deleted"), r"Items\[2\]\.Status"),
    (lambda o: o["Items"][0].pop("StationMm"), r"Items\[0\]\.StationMm"),
    (lambda o: o["Items"][0].update(StationMm=True), r"Items\[0\]\.StationMm"),
    (lambda o: o["Items"][0].update(StationMm=float("nan")), r"Items\[0\]\.StationMm"),
    (lambda o: o["Items"][1].update(TrayPointMm=[1, 2]), r"Items\[1\]\.TrayPointMm"),
    (lambda o: o["Items"][1].update(SupportId=""), r"Items\[1\]\.SupportId"),
    (lambda o: o["Items"][1].update(FittingId=5), r"Items\[1\]\.FittingId"),
    (lambda o: o["LayoutIssues"][0].pop("Code"), r"LayoutIssues\[0\]\.Code"),
    (lambda o: o.update(Notes="x"), "Notes"),
])
def test_invalid_batch_is_rejected_with_field_name(mutate, field):
    o = batch()
    mutate(o)
    with pytest.raises(HR.HangerReportError, match=field):
        HR.parse_batch(o)


def test_batch_without_items_exports_an_empty_list():
    wb = workbook(HR.build_workbook(HR.parse_batch({"BatchId": "b", "Items": []})))
    assert len(sheet_rows(wb["吊架清單"])) == 1


# ───────────── Excel 內容 ─────────────
def test_workbook_sheets_header_and_one_row_per_hanger():
    data = HR.build_workbook(HR.parse_batch(batch()))
    wb = workbook(data)
    assert wb.sheetnames == ["吊架清單", "批次摘要", "佈置問題", "參數對照"]
    cols = HR.load_columns()
    r = sheet_rows(wb["吊架清單"])
    assert r[0] == [c["label"] for c in cols]
    assert len(r) == 1 + 8
    row = dict(zip([c["key"] for c in cols], r[2]))
    assert row["seq"] == 2 and row["support_id"] == "H002" and row["tray_id"] == "101" and row["type"] == "直線段"
    assert row["status"] == "已存在" and row["status_code"] == "exists" and row["spacing_mm"] == 1500.0
    assert (row["tray_x_mm"], row["tray_y_mm"], row["tray_z_mm"]) == (2800.0, 2000.0, 3000.0)
    assert row["insert_z_mm"] == 2900.0 and row["detail"] == "本工具已放置（元素 9001）"
    elbow = dict(zip([c["key"] for c in cols], r[7]))
    assert elbow["type"] == "彎頭中點" and elbow["fitting_id"] == "201" and elbow["spacing_mm"] is None
    assert [x[cols.index(next(c for c in cols if c["key"] == "status"))] for x in r[1:]] == \
        ["新增", "已存在", "待確認", "新增", "新增（已建立）", "失敗", "新增", "新增"]
    s = dict((k, v) for k, v in sheet_rows(wb["批次摘要"])[1:] if k not in ("說明", "附註"))
    assert s["批次"] == "20261021-125357-abcdef12" and s["吊架總數"] == 8 and s["狀態：待確認（pending）"] == 1
    assert s["佈置問題中待確認（Level=pending）筆數"] == 1
    assert sheet_rows(wb["佈置問題"])[1] == ["pending", "fixed_remainder", "102", 1.0, 2.0, 3.0, "剩餘 12.000 mm"]


def test_text_that_looks_like_a_formula_stays_text_and_control_chars_are_visible():
    o = batch()
    o["Items"][0]["Detail"] = "=HYPERLINK(\"http://x\",\"y\")"
    o["Items"][1]["Detail"] = "a\u0001b"
    wb = workbook(HR.build_workbook(HR.parse_batch(o)))
    col = [c["key"] for c in HR.load_columns()].index("detail") + 1
    ws = wb["吊架清單"]
    assert ws.cell(row=2, column=col).data_type == "s"
    assert ws.cell(row=2, column=col).value == "=HYPERLINK(\"http://x\",\"y\")"
    assert ws.cell(row=3, column=col).value == "a\\u0001b"


def test_cli_writes_xlsx_and_rejects_bad_input(tmp_path, capsys):
    src = tmp_path / "cfg.json.hangers_20261021-125357-abcdef12.json"
    src.write_text("﻿" + json.dumps(batch(), ensure_ascii=False, indent=2), encoding="utf-8")   # 容忍 BOM
    assert HR.main([str(src)]) == 0
    out = tmp_path / "hangers_20261021-125357-abcdef12.xlsx"
    assert len(sheet_rows(workbook(out.read_bytes())["吊架清單"])) == 9
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps({"Items": []}), encoding="utf-8")
    assert HR.main([str(bad), "-o", str(tmp_path / "x.xlsx")]) == 2
    assert "BatchId" in capsys.readouterr().err and not (tmp_path / "x.xlsx").exists()


def test_safe_filename():
    assert HR.safe_filename('20261021-125357-ab"/../x') == "hangers_20261021-125357-ab__.._x.xlsx"


# ───────────── 網頁 ④ 區 ─────────────
@pytest.fixture
def srv(tmp_path, monkeypatch):
    monkeypatch.setenv("MEP_OUTPUT_ROOT", str(tmp_path / "out"))
    monkeypatch.setattr(acad, "find_accore", lambda: None)
    monkeypatch.setattr(X, "find_oda", lambda: None)
    s = W.make_server(0, make_dwg=False)
    threading.Thread(target=s.serve_forever, daemon=True).start()
    yield s
    s.close()


def post(s, body: bytes, origin=True, ctype="application/json", length=None):
    import http.client
    c = http.client.HTTPConnection("127.0.0.1", s.port, timeout=20)
    h = {"Host": f"127.0.0.1:{s.port}", "Content-Type": ctype}
    if length is not None:                                    # 只宣告長度、不真的送：伺服器應不讀本文直接 413
        h["Content-Length"] = str(length)
    if origin:
        h["Origin"] = f"http://127.0.0.1:{s.port}"
    c.request("POST", f"/t/{s.token}/api/hanger-report", body=body, headers=h)
    r = c.getresponse()
    out = (r.status, {k.lower(): v for k, v in r.getheaders()}, r.read())
    c.close()
    return out


def test_index_has_export_button_in_the_results_section(srv):
    from tests.test_webui import call, T
    html_ = call(srv, "GET", T(srv))[2].decode("utf-8")
    res = html_[html_.index('<section aria-labelledby="h-res">'):]
    assert 'id="hanger_report_export"' in res and ">匯出 Excel</button>" in res
    assert '<label for="hanger_report_file">' in res and html_.count("<section") == 4
    js = call(srv, "GET", T(srv, "app.js"))[2].decode("utf-8")
    assert '"api/hanger-report"' in js and "createObjectURL" in js


def test_export_endpoint_returns_xlsx(srv):
    st, h, data = post(srv, json.dumps(batch(), ensure_ascii=False).encode("utf-8"))
    assert st == 200 and h["content-type"] == HR.XLSX_TYPE
    assert h["content-disposition"] == 'attachment; filename="hangers_20261021-125357-abcdef12.xlsx"'
    assert h["x-hanger-count"] == "8" and h["content-security-policy"] == W.DOWNLOAD_CSP
    assert len(sheet_rows(workbook(data)["吊架清單"])) == 9


def test_export_endpoint_errors(srv):
    o = batch()
    o["Items"][0]["Status"] = "zzz"
    st, _, d = post(srv, json.dumps(o).encode())
    e = json.loads(d)["error"]
    assert st == 400 and e["code"] == "invalid_hanger_report" and "Items[0].Status" in e["message"]
    assert post(srv, b"{bad")[0] == 400
    assert post(srv, b"{}", origin=False)[0] == 403
    assert post(srv, b"{}", ctype="text/plain")[0] == 415
    assert post(srv, b"{}", length=W.MAX_HANGER_REPORT_BODY + 1)[0] == 413
    big = json.dumps({**batch(), "Notes": ["x" * (W.MAX_BODY + 10)]}).encode()  # 超過 /api/run 的 1 MB，但在 8 MB 內
    assert post(srv, big)[0] == 200


def test_run_endpoint_keeps_its_1mb_limit(srv):
    import http.client
    c = http.client.HTTPConnection("127.0.0.1", srv.port, timeout=20)
    c.request("POST", f"/t/{srv.token}/api/run", body=b"{}", headers={
        "Host": f"127.0.0.1:{srv.port}", "Origin": f"http://127.0.0.1:{srv.port}", "Content-Type": "application/json",
        "Content-Length": str(W.MAX_BODY + 1)})
    assert c.getresponse().status == 413
    c.close()


# ───────────── 第 2、3、5 項：吊桿命中資訊、計畫間距、欄寬 ─────────────
def test_rod_hit_audit_fields_and_type_name_are_exported():
    b = HR.parse_batch({"BatchId": "b2", "Items": [dict(
        item("H1", "1", 300.0, "new", "spacing"), HitElementId="7001", HitLinkInstanceId="8001",
        HitCategory="Structural Framing", HitZMm=3010.5, TypeName="L1200_R2500")]})
    it = b["Items"][0]
    assert (it["HitElementId"], it["HitLinkInstanceId"], it["HitCategory"], it["HitZMm"], it["TypeName"]) == \
        ("7001", "8001", "Structural Framing", 3010.5, "L1200_R2500")
    row = dict(zip([c["key"] for c in HR.load_columns()], HR.rows(b, HR.load_columns())[0]))
    assert row["rod_hit_element"] == "7001" and row["rod_hit_link"] == "8001" and row["type_name"] == "L1200_R2500"
    assert row["rod_hit_z_mm"] == 3010.5


def test_rod_hit_fields_absent_in_older_reports_export_as_blank():
    b = HR.parse_batch({"BatchId": "b3", "Items": [item("H1", "1", 300.0, "new", "spacing")]})
    row = dict(zip([c["key"] for c in HR.load_columns()], HR.rows(b, HR.load_columns())[0]))
    assert row["rod_hit_element"] is None and row["type_name"] is None and row["rod_hit_z_mm"] is None


def test_rod_hit_z_must_be_finite_number():
    with pytest.raises(HR.HangerReportError, match="HitZMm"):
        HR.parse_batch({"BatchId": "b4", "Items": [dict(item("H1", "1", 300.0, "new", "spacing"), HitZMm="3010")]})


def test_spacing_column_is_labelled_as_planned_and_summary_says_so():
    labels = [c["label"] for c in HR.load_columns() if c["key"] == "spacing_mm"]
    assert labels == ["與同橋架前一支計畫間距 (mm)"]
    summary = HR.build_workbook(HR.parse_batch({"BatchId": "b5", "Items": [item("H1", "1", 300.0, "failed", "spacing")]}))
    ws = load_workbook(io.BytesIO(summary))["批次摘要"]
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("計畫間距" in t and "失敗或待確認的吊架仍計入" in t for t in texts)


def test_column_width_counts_cjk_as_double():
    from mep_tray.hanger_report import _width
    assert _width("abc") == 3 and _width("中文") == 4 and _width("Station (mm)") == 12
    wb = load_workbook(io.BytesIO(HR.build_workbook(HR.parse_batch({"BatchId": "b6", "Items": [item("H1", "1", 300.0, "new", "spacing")]}))))
    ws = wb["吊架清單"]
    # 「與同橋架前一支計畫間距 (mm)」是中文字多的表頭，欄寬應明顯寬於只計字元數的版本
    col = [c for c in HR.load_columns()].index(next(c for c in HR.load_columns() if c["key"] == "spacing_mm")) + 1
    from openpyxl.utils import get_column_letter
    assert ws.column_dimensions[get_column_letter(col)].width >= 8
