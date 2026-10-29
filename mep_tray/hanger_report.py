"""吊架批次結果 → Excel 報表（只讀 Revit 寫出的 `<設定檔>.hangers_<批次>.json`，不碰 Revit、不改 add-in）。

輸入格式 = revit/MepTrayImport/HangerPlacer.cs 的 HangerRunReport（System.Text.Json 預設 PascalCase 鍵）。
每支吊架一列；欄位與 Revit 吊架參數名稱的對照放在 hanger_report_columns.json（revit_param 未填 = 未設定，不推測）。

衍生欄位的規則（全部只由報告內容決定，不猜）：
- 類型：Reason 為 elbow_mid／tee_mid 且有 FittingId → 彎頭中點／三通中點；Reason 為直線段佈置原因
  （spacing／near_fitting／gap_fix／short_midpoint）且沒有 FittingId → 直線段；其餘組合 → 「無法判定」並附原始碼。
- 狀態：new → 新增、created → 新增（已建立）、exists → 已存在、pending → 待確認、failed → 失敗；
  其他狀態碼拒絕（不把未知狀態歸到任何一類）。
- 間距：同一 TrayId 的直線段吊架依 StationMm 排序後與前一支的差；每支橋架的第一支與中點吊架留空。
  跨越接頭或相鄰橋架的間距不計（Station 只在同一支橋架內可比較）。
"""
from __future__ import annotations

import argparse
import io
import json
import math
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

COLUMNS_PATH = Path(__file__).with_name("hanger_report_columns.json")
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

STATUS_LABEL = {"new": "新增", "created": "新增（已建立）", "exists": "已存在", "pending": "待確認", "failed": "失敗"}
STRAIGHT_REASONS = ("spacing", "near_fitting", "gap_fix", "short_midpoint")
FITTING_REASONS = {"elbow_mid": "彎頭中點", "tee_mid": "三通中點"}
TYPE_STRAIGHT = "直線段"
TYPE_UNKNOWN = "無法判定"
DERIVED_KEYS = ("seq", "type", "spacing_mm", "status")
NUMBER_FORMAT = "0.0"

# Excel／XML 不接受的控制字元（保留 \t \n \r）；以 \\uXXXX 可見化，不靜默刪除
_XML_BAD = re.compile("[\u0000-\u0008\u000b\u000c\u000e-\u001f￾￿]")
_SOURCE_RE = re.compile(r"^([A-Za-z]+)(?:\[([0-2])\])?$")


class HangerReportError(ValueError):
    """輸入不是可匯出的吊架批次結果（訊息點名欄位）。"""


# ───────────────────────── 欄位對照表 ─────────────────────────
def load_columns(path: Path | str | None = None) -> list[dict]:
    p = Path(path) if path is not None else COLUMNS_PATH
    try:
        obj = json.loads(p.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, ValueError) as e:
        raise HangerReportError(f"無法讀取欄位對照表 {p.name}：{type(e).__name__}") from None
    return validate_columns(obj)


def validate_columns(obj) -> list[dict]:
    if not isinstance(obj, dict) or obj.get("version") != 1 or not isinstance(obj.get("columns"), list) or not obj["columns"]:
        raise HangerReportError("欄位對照表：需為 {\"version\": 1, \"columns\": [...]}（非空）")
    cols, keys, params = [], set(), set()
    for i, c in enumerate(obj["columns"]):
        f = f"columns[{i}]"
        if not isinstance(c, dict) or set(c) != {"key", "source", "label", "revit_param"}:
            raise HangerReportError(f"欄位對照表 {f}：只能且必須有 key、source、label、revit_param")
        k, src, label, rp = c["key"], c["source"], c["label"], c["revit_param"]
        if not isinstance(k, str) or not k or k in keys:
            raise HangerReportError(f"欄位對照表 {f}.key：需為不重複的非空字串")
        if not isinstance(label, str) or not label.strip():
            raise HangerReportError(f"欄位對照表 {f}.label：需為非空字串")
        if src is None:
            if k not in DERIVED_KEYS:
                raise HangerReportError(f"欄位對照表 {f}.source：null 只用於衍生欄位 {list(DERIVED_KEYS)}")
        elif not isinstance(src, str) or not _SOURCE_RE.match(src):
            raise HangerReportError(f"欄位對照表 {f}.source：需為報告欄位名稱（例如 StationMm 或 TrayPointMm[0]）")
        if rp is not None:
            if not isinstance(rp, str) or not rp or rp != rp.strip():
                raise HangerReportError(f"欄位對照表 {f}.revit_param：需為 null 或前後無空白的非空字串")
            if rp in params:
                raise HangerReportError(f"欄位對照表 {f}.revit_param：「{rp}」重複")
            params.add(rp)
        keys.add(k)
        cols.append({"key": k, "source": src, "label": label, "revit_param": rp})
    return cols


# ───────────────────────── 批次結果驗證 ─────────────────────────
def _finite(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def _req_str(it: dict, k: str, f: str, allow_empty=False) -> str:
    v = it.get(k)
    if not isinstance(v, str) or (not allow_empty and not v):
        raise HangerReportError(f"{f}.{k}：需為{'' if allow_empty else '非空'}字串")
    return v


def _opt_str(it: dict, k: str, f: str) -> str | None:
    v = it.get(k)
    if v is not None and not isinstance(v, str):
        raise HangerReportError(f"{f}.{k}：需為字串或 null")
    return v


def _req_num(it: dict, k: str, f: str) -> float:
    v = it.get(k)
    if not _finite(v):
        raise HangerReportError(f"{f}.{k}：需為有限數字")
    return float(v)


def _opt_num(it: dict, k: str, f: str) -> float | None:
    v = it.get(k)
    if v is None:
        return None
    if not _finite(v):
        raise HangerReportError(f"{f}.{k}：需為有限數字或 null")
    return float(v)


def _point(it: dict, k: str, f: str) -> list[float]:
    v = it.get(k)
    if not isinstance(v, list) or len(v) != 3 or not all(_finite(x) for x in v):
        raise HangerReportError(f"{f}.{k}：需為 3 個有限數字")
    return [float(x) for x in v]


def parse_batch(obj) -> dict:
    """驗證並正規化 HangerRunReport。缺欄位、型別錯誤、未知狀態碼一律 HangerReportError（點名欄位）。"""
    if not isinstance(obj, dict):
        raise HangerReportError("批次結果需為 JSON 物件")
    if not isinstance(obj.get("BatchId"), str) or not obj["BatchId"]:
        raise HangerReportError("BatchId：缺少（這不是吊架批次結果 *.hangers_<批次>.json？）")
    if not isinstance(obj.get("Items"), list):
        raise HangerReportError("Items：缺少或不是清單")
    items = []
    for i, it in enumerate(obj["Items"]):
        f = f"Items[{i}]"
        if not isinstance(it, dict):
            raise HangerReportError(f"{f}：需為物件")
        status = _req_str(it, "Status", f)
        if status not in STATUS_LABEL:
            raise HangerReportError(f"{f}.Status：未知的狀態碼「{status}」（只接受 {list(STATUS_LABEL)}）")
        items.append({
            "SupportId": _req_str(it, "SupportId", f), "TrayId": _req_str(it, "TrayId", f),
            "TrayUniqueId": _req_str(it, "TrayUniqueId", f, allow_empty=True), "FittingId": _opt_str(it, "FittingId", f),
            "StationMm": _req_num(it, "StationMm", f), "TrayPointMm": _point(it, "TrayPointMm", f),
            "InsertionPointMm": _point(it, "InsertionPointMm", f), "TrayDirectionDeg": _req_num(it, "TrayDirectionDeg", f),
            "RotationDeg": _req_num(it, "RotationDeg", f), "CrossMemberLengthMm": _req_num(it, "CrossMemberLengthMm", f),
            "RodLengthMm": _opt_num(it, "RodLengthMm", f), "Reason": _req_str(it, "Reason", f, allow_empty=True),
            "Status": status, "Detail": _req_str(it, "Detail", f, allow_empty=True), "ElementId": _opt_str(it, "ElementId", f),
        })
    issues = []
    for i, x in enumerate(obj.get("LayoutIssues") or []):
        f = f"LayoutIssues[{i}]"
        if not isinstance(x, dict):
            raise HangerReportError(f"{f}：需為物件")
        loc = x.get("LocationMm")
        issues.append({"Level": _req_str(x, "Level", f), "Code": _req_str(x, "Code", f),
                       "PieceId": _req_str(x, "PieceId", f, allow_empty=True),
                       "LocationMm": _point(x, "LocationMm", f) if loc is not None else None,
                       "Detail": _req_str(x, "Detail", f, allow_empty=True)})

    def strs(k):
        v = obj.get(k) or []
        if not isinstance(v, list) or not all(isinstance(s, str) for s in v):
            raise HangerReportError(f"{k}：需為字串清單")
        return v

    meta = {}
    for k in ("Entry", "RodMode", "Outcome", "Coordinates"):
        v = obj.get(k)
        if v is not None and not isinstance(v, str):
            raise HangerReportError(f"{k}：需為字串")
        meta[k] = v or ""
    return {"BatchId": obj["BatchId"], **meta, "PreflightErrors": strs("PreflightErrors"), "Notes": strs("Notes"),
            "RevitWarnings": strs("RevitWarnings"), "LayoutIssues": issues, "Items": items}


def load_batch(path: Path | str) -> dict:
    try:
        obj = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeDecodeError, ValueError) as e:
        raise HangerReportError(f"無法讀取批次結果：{type(e).__name__}") from None
    return parse_batch(obj)


# ───────────────────────── 衍生欄位 ─────────────────────────
def hanger_type(it: dict) -> str:
    r, fid = it["Reason"], it["FittingId"]
    if r in FITTING_REASONS and fid:
        return FITTING_REASONS[r]
    if r in STRAIGHT_REASONS and not fid:
        return TYPE_STRAIGHT
    return f"{TYPE_UNKNOWN}（Reason={r or '空白'}，FittingId={fid or '無'}）"


def spacings(items: list[dict]) -> list[float | None]:
    """同一 TrayId 的直線段吊架依 StationMm 排序後與前一支的差（mm）；其餘為 None。"""
    out: list[float | None] = [None] * len(items)
    groups: dict[str, list[int]] = {}
    for i, it in enumerate(items):
        if hanger_type(it) == TYPE_STRAIGHT:
            groups.setdefault(it["TrayId"], []).append(i)
    for idx in groups.values():
        idx.sort(key=lambda i: items[i]["StationMm"])
        for a, b in zip(idx, idx[1:]):
            out[b] = round(items[b]["StationMm"] - items[a]["StationMm"], 3)
    return out


def rows(batch: dict, columns: list[dict]) -> list[list]:
    items = batch["Items"]
    sp = spacings(items)
    out = []
    for n, it in enumerate(items):
        derived = {"seq": n + 1, "type": hanger_type(it), "spacing_mm": sp[n], "status": STATUS_LABEL[it["Status"]]}
        row = []
        for c in columns:
            if c["source"] is None:
                row.append(derived[c["key"]])
                continue
            name, idx = _SOURCE_RE.match(c["source"]).groups()
            if name not in it:
                raise HangerReportError(f"欄位對照表 {c['key']}.source：報告沒有欄位 {name}")
            v = it[name]
            if idx is not None:
                if not isinstance(v, list):
                    raise HangerReportError(f"欄位對照表 {c['key']}.source：{name} 不是陣列")
                v = v[int(idx)]
            row.append(v)
        out.append(row)
    return out


def counts(batch: dict) -> dict:
    c = {k: 0 for k in STATUS_LABEL}
    for it in batch["Items"]:
        c[it["Status"]] += 1
    return c


# ───────────────────────── Excel ─────────────────────────
def _cell_text(v):
    """字串一律當文字（不讓 = + - @ 開頭的值變成公式），控制字元可見化。"""
    if isinstance(v, str):
        return _XML_BAD.sub(lambda m: "\\u%04X" % ord(m.group()), v)
    return v


def _write(ws, r: int, c: int, v):
    cell = ws.cell(row=r, column=c, value=_cell_text(v))
    if isinstance(v, str):
        cell.data_type = "s"
    elif isinstance(v, float):
        cell.number_format = NUMBER_FORMAT
    return cell


def _table(ws, header: list[str], body: list[list]):
    bold = Font(bold=True)
    for j, h in enumerate(header, 1):
        _write(ws, 1, j, h).font = bold
    for i, row in enumerate(body, 2):
        for j, v in enumerate(row, 1):
            _write(ws, i, j, v)
    for j, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(r[j - 1])) for r in body[:500] if r[j - 1] is not None])
        ws.column_dimensions[get_column_letter(j)].width = min(max(8, width * 1.3 + 2), 60)
    ws.freeze_panes = "A2"


def build_workbook(batch: dict, columns: list[dict] | None = None) -> bytes:
    columns = columns if columns is not None else load_columns()
    wb = Workbook()
    ws = wb.active
    ws.title = "吊架清單"
    _table(ws, [c["label"] for c in columns], rows(batch, columns))
    if batch["Items"]:
        ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("批次摘要")
    c = counts(batch)
    summary = [["批次", batch["BatchId"]], ["結果", batch["Outcome"]], ["入口", batch["Entry"]], ["吊桿模式", batch["RodMode"]],
               ["座標", batch["Coordinates"]], ["吊架總數", len(batch["Items"])]]
    summary += [[f"狀態：{STATUS_LABEL[k]}（{k}）", c[k]] for k in STATUS_LABEL]
    summary += [["佈置問題（LayoutIssues）筆數", len(batch["LayoutIssues"])],
                ["佈置問題中待確認（Level=pending）筆數", sum(1 for x in batch["LayoutIssues"] if x["Level"] == "pending")]]
    summary += [["預檢錯誤", e] for e in batch["PreflightErrors"]]
    summary += [["附註", n] for n in batch["Notes"]]
    summary += [["Revit 警告", w] for w in batch["RevitWarnings"]]
    summary += [["說明", "本報表只讀取批次結果 JSON 產生，未連線 Revit；「新增」在預覽／已取消的批次中表示計畫新增，"
                         "實際是否建立以「原始狀態碼」created 與「已建立吊架 ElementId」為準。"],
                ["說明", "間距＝同一橋架（TrayId）直線段吊架依 Station 排序後與前一支的差；跨接頭或跨橋架不計，"
                         "每支橋架第一支與中點吊架留空。"]]
    _table(s, ["項目", "內容"], summary)

    li = wb.create_sheet("佈置問題")
    _table(li, ["等級", "代碼", "橋架／接頭 ElementId", "X (mm)", "Y (mm)", "Z (mm)", "說明"],
           [[x["Level"], x["Code"], x["PieceId"], *(x["LocationMm"] or [None, None, None]), x["Detail"]]
            for x in batch["LayoutIssues"]])

    m = wb.create_sheet("參數對照")
    _table(m, ["報表欄位", "欄位代碼", "來源（批次 JSON）", "Revit 吊架參數名稱"],
           [[c["label"], c["key"], c["source"] or "（衍生）", c["revit_param"] or "未設定"] for c in columns])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def safe_filename(batch_id: str) -> str:
    return "hangers_" + re.sub(r"[^A-Za-z0-9._-]", "_", batch_id)[:64] + ".xlsx"


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="吊架批次結果 JSON（*.hangers_<批次>.json）→ Excel 報表")
    ap.add_argument("batch_json", type=Path)
    ap.add_argument("-o", "--output", type=Path, default=None, help="輸出 .xlsx（預設：與輸入同資料夾 hangers_<批次>.xlsx）")
    ap.add_argument("--columns", type=Path, default=None, help="欄位對照表（預設：mep_tray/hanger_report_columns.json）")
    a = ap.parse_args(argv)
    try:
        batch = load_batch(a.batch_json)
        data = build_workbook(batch, load_columns(a.columns))
    except HangerReportError as e:
        print(f"錯誤：{e}", file=sys.stderr)
        return 2
    out = a.output or a.batch_json.with_name(safe_filename(batch["BatchId"]))
    out.write_bytes(data)
    print(f"已匯出 {len(batch['Items'])} 筆吊架：{out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
